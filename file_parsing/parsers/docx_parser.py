from pathlib import Path

from docx import Document

from file_parsing.registry import register
from file_parsing.types import ParsedResult


@register("docx")
def parse_docx(path: Path) -> ParsedResult:
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return ParsedResult(markdown="\n".join(lines), status="success", parser_used="python_docx")
