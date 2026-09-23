import os
import shutil
import subprocess
from pathlib import Path

from docx import Document


def _find_pandoc() -> str:
    configured = os.environ.get("PANDOC_PATH")
    if configured:
        return configured
    found = shutil.which("pandoc")
    if found:
        return found
    raise RuntimeError(
        "找不到 Pandoc 的可执行文件，"
        "请设置环境变量 PANDOC_PATH 指向 pandoc(.exe) 路径，或将其加入 PATH"
    )


def _extract_with_python_docx(path: Path) -> str:
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def convert_docx_to_markdown(path: Path) -> tuple[str, str]:
    try:
        pandoc_exe = _find_pandoc()
        result = subprocess.run(
            [pandoc_exe, str(path), "-t", "gfm", "--wrap=none"],
            capture_output=True,
            check=True,
        )
        return result.stdout.decode("utf-8"), "pandoc"
    except (RuntimeError, subprocess.CalledProcessError, FileNotFoundError):
        # FileNotFoundError: PANDOC_PATH 指向一个不存在的路径时，subprocess.run 直接抛这个
        # （_find_pandoc 只保证字符串非空，不校验路径真实存在）
        return _extract_with_python_docx(path), "python_docx_fallback"
