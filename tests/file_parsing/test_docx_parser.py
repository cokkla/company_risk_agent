import subprocess

from docx import Document

from file_parsing.parsers.docx_parser import parse_docx


def _make_docx(path):
    doc = Document()
    doc.add_paragraph("hello world")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "a"
    table.rows[0].cells[1].text = "b"
    doc.save(path)
    return path


def test_parse_docx_via_pandoc(tmp_path, monkeypatch):
    path = _make_docx(tmp_path / "sample.docx")

    def fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, 0, stdout=b"hello world\n\na | b\n")

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: "pandoc")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    result = parse_docx(path)

    assert result.status == "success"
    assert result.parser_used == "pandoc"
    assert "hello world" in result.markdown
    assert "a | b" in result.markdown


def test_parse_docx_falls_back_to_python_docx_when_pandoc_missing(tmp_path, monkeypatch):
    path = _make_docx(tmp_path / "sample.docx")

    def raise_not_found():
        raise RuntimeError("pandoc not found")

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", raise_not_found)

    result = parse_docx(path)

    assert result.status == "success"
    assert result.parser_used == "python_docx_fallback"
    assert "hello world" in result.markdown
    assert "a | b" in result.markdown


def test_parse_docx_falls_back_when_pandoc_exits_nonzero(tmp_path, monkeypatch):
    path = _make_docx(tmp_path / "sample.docx")

    def fake_run(cmd, **kwargs):
        raise subprocess.CalledProcessError(1, cmd)

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: "pandoc")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    result = parse_docx(path)

    assert result.status == "success"
    assert result.parser_used == "python_docx_fallback"
    assert "hello world" in result.markdown
    assert "a | b" in result.markdown


def test_parse_docx_raises_when_both_pandoc_and_fallback_fail(tmp_path, monkeypatch):
    missing_path = tmp_path / "does_not_exist.docx"

    def fake_run(cmd, **kwargs):
        raise subprocess.CalledProcessError(1, cmd)

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: "pandoc")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    try:
        parse_docx(missing_path)
        assert False, "应该抛出异常"
    except Exception:
        pass
