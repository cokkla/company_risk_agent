import openpyxl

from file_parsing.parsers.xlsx_parser import parse_xlsx


def test_parse_xlsx_basic_content(tmp_path):
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["a", "b"])
    ws.append([1, 2])
    wb.save(path)

    result = parse_xlsx(path)

    assert result.status == "success"
    assert "## Sheet1" in result.markdown
    assert "a | b" in result.markdown
    assert "1 | 2" in result.markdown


def test_parse_xlsx_escapes_pipe_in_cell_value(tmp_path):
    # 回归测试：单元格内容里的 "|" 若不转义，会被下游 markdown 解析成多出的列
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a|b", "normal"])
    wb.save(path)

    result = parse_xlsx(path)

    assert "a\\|b | normal" in result.markdown
    assert "a|b | normal" not in result.markdown


def test_parse_xlsx_none_cell_becomes_empty_string(tmp_path):
    # 尾部的空单元格会被 openpyxl 直接裁掉，因此把 None 放在中间以验证渲染为空字符串
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", None, "c"])
    wb.save(path)

    result = parse_xlsx(path)

    assert "a |  | c" in result.markdown
