import subprocess

from docx import Document

from file_parsing.parsers._docx_pandoc import _find_pandoc, convert_docx_to_markdown


def test_find_pandoc_uses_env_var(monkeypatch):
    monkeypatch.setenv("PANDOC_PATH", r"C:\Users\30684\AppData\Local\Pandoc\pandoc.exe")
    assert _find_pandoc() == r"C:\Users\30684\AppData\Local\Pandoc\pandoc.exe"


def test_find_pandoc_falls_back_to_path(monkeypatch):
    monkeypatch.delenv("PANDOC_PATH", raising=False)
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.shutil.which", lambda name: "/usr/bin/pandoc")
    assert _find_pandoc() == "/usr/bin/pandoc"


def test_find_pandoc_raises_when_not_found(monkeypatch):
    monkeypatch.delenv("PANDOC_PATH", raising=False)
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.shutil.which", lambda name: None)
    try:
        _find_pandoc()
        assert False, "应该抛出 RuntimeError"
    except RuntimeError:
        pass


def _make_docx(path):
    doc = Document()
    doc.add_paragraph("hello world")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "a"
    table.rows[0].cells[1].text = "b"
    doc.save(path)
    return path


def test_convert_docx_to_markdown_via_pandoc(tmp_path, monkeypatch):
    path = _make_docx(tmp_path / "sample.docx")

    def fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, 0, stdout=b"# fake markdown\n")

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: "pandoc")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    markdown, parser_used = convert_docx_to_markdown(path)

    assert parser_used == "pandoc"
    assert markdown == "# fake markdown\n"


def test_convert_docx_to_markdown_falls_back_when_pandoc_not_found(tmp_path, monkeypatch):
    path = _make_docx(tmp_path / "sample.docx")

    def raise_not_found():
        raise RuntimeError("pandoc not found")

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", raise_not_found)

    markdown, parser_used = convert_docx_to_markdown(path)

    assert parser_used == "python_docx_fallback"
    assert "hello world" in markdown
    assert "a | b" in markdown


def test_convert_docx_to_markdown_falls_back_when_pandoc_exits_nonzero(tmp_path, monkeypatch):
    path = _make_docx(tmp_path / "sample.docx")

    def fake_run(cmd, **kwargs):
        raise subprocess.CalledProcessError(1, cmd)

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: "pandoc")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    markdown, parser_used = convert_docx_to_markdown(path)

    assert parser_used == "python_docx_fallback"
    assert "hello world" in markdown
    assert "a | b" in markdown


def test_convert_docx_to_markdown_falls_back_when_pandoc_path_is_invalid(tmp_path, monkeypatch):
    # 回归测试：PANDOC_PATH 指向不存在的路径时，subprocess.run 直接抛 FileNotFoundError，
    # 这个异常类型之前没被捕获，导致本该降级的场景直接把异常甩了出去。
    path = _make_docx(tmp_path / "sample.docx")

    def fake_run(cmd, **kwargs):
        raise FileNotFoundError(2, "系统找不到指定的文件")

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: r"C:\not\a\real\path\pandoc.exe")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    markdown, parser_used = convert_docx_to_markdown(path)

    assert parser_used == "python_docx_fallback"
    assert "hello world" in markdown
    assert "a | b" in markdown


def test_convert_docx_to_markdown_raises_when_both_fail(tmp_path, monkeypatch):
    missing_path = tmp_path / "does_not_exist.docx"

    def fake_run(cmd, **kwargs):
        raise subprocess.CalledProcessError(1, cmd)

    monkeypatch.setattr("file_parsing.parsers._docx_pandoc._find_pandoc", lambda: "pandoc")
    monkeypatch.setattr("file_parsing.parsers._docx_pandoc.subprocess.run", fake_run)

    try:
        convert_docx_to_markdown(missing_path)
        assert False, "应该抛出异常"
    except Exception:
        pass
